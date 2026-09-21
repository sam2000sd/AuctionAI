
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse, parse_qs, unquote
import os
import re
import sys
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

# Streamlit Cloud runs this file from inside /app, so add the repo root
# to Python path before importing app.core/app.scrapers modules.
ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from app.core.config import SCRAPED_DIR, EXPORT_DIR
from app.core.normalize import normalize_files
from app.core.formulas import calc_bid
from app.core.utils import money, pct, this_or_next_week, city_from_address, address_key, zillow_link, redfin_link, redfin_search_link, loose_address_key, MD_COUNTIES
from app.scrapers.sources import scrape_many, clear_cache, load_values
from app.storage.local import load_bids, save_bids, merge_bids, load_hidden, hide_address, clear_hidden, load_blocked_cities, save_blocked_cities, load_favorite_properties, save_favorite_properties, toggle_favorite_property, load_layout_defaults, save_layout_defaults, remote_enabled

st.set_page_config(page_title="Auction Intelligence", page_icon="🏛️", layout="wide")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

:root {
  --ai-bg: #f5f7fb;
  --ai-card: #ffffff;
  --ai-ink: #0f172a;
  --ai-ink-2: #334155;
  --ai-muted: #64748b;
  --ai-line: #e2e8f0;
  --ai-line-2: #cbd5e1;
  --ai-navy: #1e3a5f;
  --ai-navy-2: #16304f;
  --ai-accent: #2563eb;
  --ai-sidebar: #0f1b2d;
  --ai-sidebar-2: #152a44;
  --ai-radius: 10px;
}

html, body, [data-testid="stAppViewContainer"], .stApp {
  font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif !important;
  background: var(--ai-bg) !important;
  color: var(--ai-ink);
}
[data-testid="stAppViewContainer"] *:not([data-testid="stIconMaterial"]):not(.material-symbols-rounded) { font-family: 'Inter', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; }
[data-testid="stIconMaterial"], .material-symbols-rounded { font-family: 'Material Symbols Rounded' !important; }

/* Hide Streamlit chrome (toolbar, hamburger, footer) for a product look */
[data-testid="stToolbar"], #MainMenu, footer, [data-testid="stDecoration"] { display: none !important; }
header[data-testid="stHeader"] { background: transparent !important; height: 0 !important; }
.block-container { max-width: 1900px; padding: 1.4rem 2rem 3rem 2rem; }

/* ---------- Sidebar ---------- */
[data-testid="stSidebar"] { background: linear-gradient(180deg, var(--ai-sidebar) 0%, var(--ai-sidebar-2) 100%) !important; border-right: 1px solid rgba(255,255,255,.06); }
[data-testid="stSidebar"] > div:first-child { padding-top: 1.2rem; }
[data-testid="stSidebar"] h1, [data-testid="stSidebar"] h2, [data-testid="stSidebar"] h3 {
  color: #94a3b8 !important; font-size: .7rem !important; font-weight: 700 !important;
  letter-spacing: .12em; text-transform: uppercase; margin: .2rem 0 .6rem 0 !important; padding: 0 !important;
}
[data-testid="stSidebar"] label, [data-testid="stSidebar"] label p, [data-testid="stSidebar"] label span,
[data-testid="stSidebar"] p, [data-testid="stSidebar"] span { color: #e2e8f0 !important; -webkit-text-fill-color: #e2e8f0 !important; opacity: 1 !important; }
[data-testid="stSidebar"] label, [data-testid="stSidebar"] label p { font-size: .82rem !important; font-weight: 500 !important; }
[data-testid="stSidebar"] .stCaptionContainer, [data-testid="stSidebar"] .stCaptionContainer p { color: #7c8ba1 !important; -webkit-text-fill-color: #7c8ba1 !important; font-size: .74rem !important; }
[data-testid="stSidebar"] hr { border-color: rgba(255,255,255,.08); margin: 1.1rem 0; }
[data-testid="stSidebar"] div[data-baseweb="select"] > div,
[data-testid="stSidebar"] textarea,
[data-testid="stSidebar"] input {
  background: rgba(255,255,255,.06) !important; color: #f8fafc !important;
  border: 1px solid rgba(255,255,255,.12) !important; border-radius: 8px !important;
}
[data-testid="stSidebar"] div[data-baseweb="select"] svg { fill: #cbd5e1 !important; }
[data-testid="stSidebar"] div[data-baseweb="tag"] { background: rgba(255,255,255,.14) !important; border-radius: 6px !important; }
[data-testid="stSidebar"] div[data-baseweb="tag"] span { color: #f8fafc !important; -webkit-text-fill-color: #f8fafc !important; font-weight: 600; font-size: .74rem; }
[data-testid="stSidebar"] div.stButton > button {
  background: rgba(255,255,255,.06) !important; color: #e2e8f0 !important;
  border: 1px solid rgba(255,255,255,.14) !important; border-radius: 8px !important;
  font-weight: 600 !important; font-size: .8rem !important; min-height: 38px; padding: 0 6px !important;
}
[data-testid="stSidebar"] div.stButton > button:hover { background: rgba(255,255,255,.12) !important; border-color: rgba(255,255,255,.28) !important; }
[data-testid="stSidebar"] div.stButton > button[kind="primary"] { background: var(--ai-accent) !important; color: #fff !important; border: 0 !important; }
[data-testid="stSidebar"] div.stButton > button[kind="primary"]:hover { background: #1d4ed8 !important; }
[data-testid="stSidebar"] div[data-testid="stNumberInput"] input { background: #ffffff !important; color: var(--ai-ink) !important; -webkit-text-fill-color: var(--ai-ink) !important; }
[data-testid="stSidebar"] div[data-testid="stNumberInput"] button { background: #ffffff !important; color: var(--ai-ink) !important; border-color: #d1d5db !important; }
[data-testid="stSidebar"] div[data-testid="stNumberInput"] button svg { fill: var(--ai-ink) !important; }
[data-testid="stSidebar"] .stAlert { border-radius: 8px; }

/* ---------- Header ---------- */
.app-brand { display: flex; align-items: center; gap: 10px; margin: 0 0 6px 0; }
.app-brand .badge { background: var(--ai-navy); color: #fff; font-size: .66rem; font-weight: 800; letter-spacing: .12em; padding: 5px 11px; border-radius: 999px; }
.app-brand .env { color: var(--ai-muted); font-size: .72rem; font-weight: 500; }
h1 { font-size: 1.85rem !important; font-weight: 800 !important; letter-spacing: -0.03em; color: var(--ai-ink) !important; margin: 0 !important; padding: 0 !important; line-height: 1.15 !important; }
[data-testid="stCaptionContainer"] p { color: var(--ai-muted); font-size: .82rem; }

/* Section labels (st.subheader) */
h3 { font-size: .72rem !important; font-weight: 700 !important; letter-spacing: .12em; text-transform: uppercase; color: var(--ai-muted) !important; margin: 1.4rem 0 .5rem 0 !important; padding: 0 !important; }
h2 { font-size: 1.05rem !important; font-weight: 700 !important; color: var(--ai-ink) !important; }

/* Cards (st.container(border=True)) */
[data-testid="stVerticalBlockBorderWrapper"] { background: var(--ai-card); border: 1px solid var(--ai-line) !important; border-radius: var(--ai-radius) !important; box-shadow: 0 1px 2px rgba(15,23,42,.04), 0 6px 18px -12px rgba(15,23,42,.18); }
[data-testid="stVerticalBlockBorderWrapper"] > div { padding: .9rem 1.1rem !important; }

/* ---------- Inputs ---------- */
div[data-testid="stTextInput"] input, div[data-testid="stNumberInput"] input, div[data-baseweb="select"] > div, div[data-baseweb="base-input"] {
  background: #ffffff !important; border: 1px solid var(--ai-line-2) !important; border-radius: 8px !important; box-shadow: none !important; color: var(--ai-ink) !important; font-size: .88rem !important;
}
div[data-testid="stTextInput"] input:focus, div[data-baseweb="select"] > div:focus-within { border-color: var(--ai-accent) !important; box-shadow: 0 0 0 3px rgba(37,99,235,.14) !important; }
div[data-baseweb="tag"] { background: #e8eef7 !important; border-radius: 6px !important; }
div[data-baseweb="tag"] span { color: var(--ai-navy) !important; font-weight: 600; font-size: .76rem; }
div[data-baseweb="tag"] svg { fill: var(--ai-navy) !important; }
label[data-testid="stWidgetLabel"] p { font-size: .72rem !important; font-weight: 600 !important; letter-spacing: .04em; text-transform: uppercase; color: var(--ai-muted) !important; }
[data-testid="stRadio"] label p { text-transform: none; letter-spacing: 0; font-size: .86rem !important; font-weight: 500 !important; color: var(--ai-ink-2) !important; }
[data-testid="stRadio"] [data-testid="stWidgetLabel"] p { font-size: .72rem !important; font-weight: 600 !important; letter-spacing: .04em; text-transform: uppercase; color: var(--ai-muted) !important; }
[data-testid="stExpander"] { border: 1px solid var(--ai-line) !important; border-radius: 8px !important; background: #fff; }
[data-testid="stExpander"] summary { font-size: .86rem; font-weight: 600; color: var(--ai-ink-2); }

/* ---------- Buttons ---------- */
div.stButton > button, div.stDownloadButton > button {
  border-radius: 8px !important; font-weight: 600 !important; font-size: .86rem !important; min-height: 40px;
  border: 1px solid var(--ai-line-2) !important; background: #ffffff !important; color: var(--ai-ink) !important; box-shadow: 0 1px 2px rgba(15,23,42,.05) !important; transition: all .12s ease;
}
div.stButton > button:hover, div.stDownloadButton > button:hover { border-color: var(--ai-navy) !important; color: var(--ai-navy) !important; background: #f8fafc !important; }
div.stButton > button[kind="primary"] { background: var(--ai-navy) !important; color: #ffffff !important; border: 1px solid var(--ai-navy) !important; }
div.stButton > button[kind="primary"]:hover { background: var(--ai-navy-2) !important; color: #fff !important; }
div.stButton > button:focus:not(:active) { box-shadow: 0 0 0 3px rgba(37,99,235,.18) !important; }
a { color: var(--ai-accent); font-weight: 600; text-decoration: none; }
a:hover { text-decoration: underline; }

/* ---------- Auction grid ---------- */
.day-header { display: flex; align-items: center; justify-content: space-between; background: var(--ai-navy); color: #fff; padding: 10px 14px; border-radius: 8px; margin: 2px 0 8px 0; }
.day-header .d { font-weight: 700; font-size: .92rem; letter-spacing: -0.01em; }
.day-header .n { font-size: .72rem; font-weight: 600; color: #cbd5e1; background: rgba(255,255,255,.1); padding: 3px 9px; border-radius: 999px; }
[class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] { align-items: center !important; gap: .55rem !important; padding: 5px 0; border-bottom: 1px solid #eef2f7; }
[class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"]:last-child { border-bottom: 0; }
[class*="st-key-grid_card_"] .gh { font-size: .62rem; font-weight: 700; letter-spacing: .03em; text-transform: uppercase; color: var(--ai-muted); white-space: nowrap; overflow: visible; padding: 2px 0 6px 0; }
[class*="st-key-grid_card_"] .gc { font-size: .84rem; color: var(--ai-ink); line-height: 1.3; }
[class*="st-key-grid_card_"] .gc.mono { font-variant-numeric: tabular-nums; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
[class*="st-key-grid_card_"] .gc.county { word-break: keep-all; overflow-wrap: normal; }
[class*="st-key-grid_card_"] .gc.time { font-size: .8rem; font-variant-numeric: tabular-nums; }
[class*="st-key-grid_card_"] .gc.muted { color: var(--ai-muted); }
[class*="st-key-grid_card_"] .gc.addr { font-weight: 600; }
[class*="st-key-grid_card_"] .gc.strong { font-weight: 700; color: var(--ai-ink); }
.src { display: inline-block; white-space: nowrap; font-size: .66rem; font-weight: 800; letter-spacing: .06em; padding: 3px 7px; border-radius: 6px; background: #e8eef7; color: var(--ai-navy); }
.src.AC { background: #e0f2fe; color: #075985; }
.src.TW { background: #dcfce7; color: #166534; }
.src.HW { background: #fef3c7; color: #92400e; }
.src.MWC { background: #f3e8ff; color: #6b21a8; }
.src.BL { background: #ffe4e6; color: #9f1239; }
.src.ADC { background: #e2e8f0; color: #1e293b; }
[class*="st-key-grid_card_"] .pill { display: inline-block; font-size: .74rem; font-weight: 700; padding: 5px 10px; border-radius: 6px; background: #eef2f7; color: var(--ai-navy); text-decoration: none; white-space: nowrap; }
[class*="st-key-grid_card_"] .pill:hover { background: var(--ai-navy); color: #fff; text-decoration: none; }
[class*="st-key-grid_card_"] .links { white-space: nowrap; }
[class*="st-key-grid_card_"] .links a { display: inline-block; font-size: .72rem; font-weight: 800; padding: 4px 7px; border-radius: 6px; background: #eef2f7; color: var(--ai-navy); margin-right: 4px; }
[class*="st-key-grid_card_"] .links a:hover { background: var(--ai-navy); color: #fff; text-decoration: none; }
[class*="st-key-grid_card_"] div.stButton > button { min-height: 34px; padding: 0 6px; font-size: .78rem; white-space: nowrap; }
[class*="st-key-grid_card_"] div[data-testid="stTextInput"] input { min-height: 34px; height: 34px; padding: 0 8px; font-size: .84rem !important; }
[class*="st-key-grid_card_"] div[data-baseweb="select"] > div { min-height: 34px; font-size: .82rem !important; }
[class*="st-key-grid_card_"] [data-testid="stCheckbox"] { margin: 0; }
[class*="st-key-grid_card_"] [data-testid="stCheckbox"] label { padding: 0; min-height: 0; }
[class*="st-key-grid_card_"] .element-container { margin-bottom: 0 !important; }

/* Tabs */
[data-testid="stTabs"] button { font-weight: 600; font-size: .86rem; }

/* ---------- Phone field-mode cards ---------- */
[class*="st-key-pcard_"] { margin-bottom: 10px; }
[class*="st-key-pcard_"] > div, [class*="st-key-pcard_"] { padding: .55rem .7rem !important; }
[class*="st-key-pcard_"] [data-testid="stHorizontalBlock"]:first-of-type { align-items: start !important; }
[class*="st-key-pcard_"] [data-testid="stHorizontalBlock"] { gap: .5rem !important; align-items: end !important; flex-wrap: nowrap !important; flex-direction: row !important; }
[class*="st-key-pcard_"] [data-testid="stHorizontalBlock"] > div { flex: 1 1 0 !important; width: auto !important; min-width: 0 !important; }
[class*="st-key-pcard_"] [data-testid="stHorizontalBlock"] > div:has(> div > div > .stButton) { flex: 0 0 52px !important; }
[class*="st-key-pcard_"] div.stButton > button { min-height: 44px; }
[class*="st-key-pcard_"] .pc-head { padding-bottom: 4px; }
[class*="st-key-pcard_"] .element-container { margin-bottom: 0 !important; }
[class*="st-key-pcard_"] label[data-testid="stWidgetLabel"] p { font-size: .62rem !important; }
[class*="st-key-pcard_"] div[data-testid="stTextInput"] input { min-height: 40px; font-size: 16px !important; }
.pc-top { display: flex; align-items: center; gap: 8px; flex-wrap: wrap; }
.pc-time { font-size: 1.05rem; font-weight: 800; color: var(--ai-navy); font-variant-numeric: tabular-nums; }
.pc-county { font-size: .76rem; color: var(--ai-muted); font-weight: 600; }
.pc-addr { font-size: 1.02rem; font-weight: 700; color: var(--ai-ink); line-height: 1.25; margin: 4px 0 2px 0; }
.pc-meta { font-size: .8rem; color: var(--ai-muted); }
.pc-meta b { color: var(--ai-ink-2); font-weight: 700; }
.pc-val { font-size: .84rem; color: var(--ai-ink); margin-top: 3px; }
.pc-val b { font-size: .98rem; color: var(--ai-navy); font-weight: 800; }
.pc-val .vl { font-size: .68rem; font-weight: 700; letter-spacing: .06em; text-transform: uppercase; color: var(--ai-muted); }
.pc-val .more { color: var(--ai-muted); font-size: .76rem; }
.gc.val { display: flex; flex-direction: column; line-height: 1.1; cursor: help; }
.gc.val .vh { font-weight: 700; color: var(--ai-navy); }
.gc.val .vl { font-size: .58rem; font-weight: 700; letter-spacing: .06em; text-transform: uppercase; color: var(--ai-muted); }
.pill.muted { background: #f1f5f9; color: #94a3b8; }
.pc-bids { display: flex; gap: 8px; margin: 8px 0 4px 0; }
.pc-bids > div { flex: 1; background: #f1f5f9; border-radius: 8px; padding: 6px 8px; display: flex; flex-direction: column; }
.pc-bids .k { font-size: .6rem; font-weight: 700; letter-spacing: .08em; color: var(--ai-muted); }
.pc-bids .v { font-size: 1.02rem; font-weight: 800; color: var(--ai-ink); font-variant-numeric: tabular-nums; }
.pc-bids > div:first-child { background: var(--ai-navy); }
.pc-bids > div:first-child .k { color: #cbd5e1; }
.pc-bids > div:first-child .v { color: #fff; }
.pill { display: inline-block; font-size: .74rem; font-weight: 700; padding: 5px 10px; border-radius: 6px; background: #eef2f7; color: var(--ai-navy); text-decoration: none; white-space: nowrap; }
.pill:hover { background: var(--ai-navy); color: #fff; text-decoration: none; }
.pill.wide { display: block; text-align: center; padding: 11px 6px; font-size: .84rem; }
.day-header.phone { position: sticky; top: 0; z-index: 5; margin: 10px 0 8px 0; }
[class*="st-key-ad_"] div.stButton > button { background: #eef2f7 !important; color: var(--ai-navy) !important; border: 0 !important; font-weight: 700 !important; padding: 0 4px !important; min-width: 0 !important; }
[class*="st-key-ad_"] div.stButton > button:hover { background: var(--ai-navy) !important; color: #fff !important; }
[data-testid="stPopover"] button { border-radius: 8px !important; min-height: 34px; padding: 0 10px !important; background: #eef2f7 !important; border: 0 !important; }
[data-testid="stPopover"] button p, [data-testid="stPopover"] button div, [data-testid="stPopover"] button span { color: var(--ai-navy) !important; font-weight: 700 !important; font-size: .78rem !important; }
[data-testid="stPopover"] button svg { fill: var(--ai-navy) !important; }
[data-testid="stPopoverBody"] { max-width: min(92vw, 720px); }
.ad-body { font-size: .82rem; line-height: 1.45; color: var(--ai-ink-2); max-height: 70vh; overflow: auto; }
.ad-body table { max-width: 100%; }

/* ---------- Phone / narrow ---------- */
@media (max-width: 900px) {
  .block-container { padding: .8rem .7rem 2rem .7rem !important; }
  h1 { font-size: 1.35rem !important; }
  [data-testid="stVerticalBlockBorderWrapper"] > div { padding: .7rem .7rem !important; }
  /* Keep the grid as a swipeable table instead of stacking 16 columns vertically */
  /* One shared horizontal scroller per day card (rows + header move together) */
  [class*="st-key-grid_card_"] { overflow-x: auto !important; -webkit-overflow-scrolling: touch; padding-bottom: 6px; }
  [class*="st-key-grid_card_"] [data-testid="stLayoutWrapper"] { width: max-content !important; min-width: 100% !important; }
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] { flex-wrap: nowrap !important; overflow: visible !important; width: max-content !important; min-width: 100% !important; align-items: center !important; }
  [class*="st-key-grid_card_"] .day-header { position: sticky; left: 0; }
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] > div { flex-shrink: 0 !important; min-width: 96px !important; }
  /* Phone column order: star, Address, Time, Src, County, then the rest */
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] > div { order: 10; }
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] > div:nth-child(1) { order: 1; min-width: 48px !important; }
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] > div:nth-child(5) { order: 2; min-width: 210px !important; }
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] > div:nth-child(2) { order: 3; min-width: 72px !important; }
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] > div:nth-child(3) { order: 4; min-width: 52px !important; }
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] > div:nth-child(4) { order: 5; min-width: 120px !important; }
  [class*="st-key-grid_card_"] [data-testid="stHorizontalBlock"] > div:nth-child(7) { min-width: 52px !important; }
  [class*="st-key-grid_card_"] .gc.addr { font-size: .82rem; }
  div.stButton > button, div.stDownloadButton > button { min-height: 44px !important; }
  input, select, textarea { font-size: 16px !important; } /* stops iOS auto-zoom */
}
</style>
""", unsafe_allow_html=True)

SCRAPED_DIR.mkdir(parents=True, exist_ok=True)
EXPORT_DIR.mkdir(parents=True, exist_ok=True)
layout_defaults = load_layout_defaults()
AUCTION_HOUSES = ["AC", "TW", "HW", "MWC", "BL", "ADC"]
SOURCE_NAMES = {"AC": "Alex Cooper", "TW": "Tidewater", "HW": "Harvey West", "MWC": "McCabe Weisberg", "BL": "A. J. Billig", "ADC": "Auction.com"}
DEFAULT_COUNTY_FILTER = ["Montgomery County", "Prince George's County", "Howard County", "Frederick County", "Anne Arundel County", "Washington, DC"]

def default_value(key, fallback):
    return layout_defaults.get(key, fallback)


def start_shutdown_listener():
    """Close local Python/Streamlit after the browser tab is really gone."""
    if st.session_state.get("shutdown_listener_started"):
        return
    st.session_state.shutdown_listener_started = True
    state = {"last_ping": datetime.now().timestamp(), "closed_at": None}

    class ShutdownHandler(BaseHTTPRequestHandler):
        def _headers(self):
            self.send_response(204)
            self.send_header("Access-Control-Allow-Origin", "*")
            self.send_header("Access-Control-Allow-Methods", "POST, OPTIONS")
            self.end_headers()

        def do_OPTIONS(self):
            self._headers()

        def do_POST(self):
            path = (self.path or "").lower()
            now = datetime.now().timestamp()
            if "ping" in path:
                state["last_ping"] = now
                state["closed_at"] = None
            elif "closed" in path or "shutdown" in path:
                state["closed_at"] = now
            self._headers()

        def log_message(self, *args):
            return

    def monitor():
        import time
        while True:
            time.sleep(1.0)
            now = datetime.now().timestamp()
            closed_at = state.get("closed_at")
            if closed_at and now - max(closed_at, state.get("last_ping", 0)) > 6:
                os._exit(0)

    def run_server():
        try:
            threading.Thread(target=monitor, daemon=True).start()
            HTTPServer(("127.0.0.1", 8765), ShutdownHandler).serve_forever()
        except OSError:
            pass

    threading.Thread(target=run_server, daemon=True).start()

if os.getenv("STREAMLIT_SERVER_HEADLESS") != "true":
    start_shutdown_listener()
components.html("""
<script>
(function () {
  function ping() {
    try { navigator.sendBeacon('http://127.0.0.1:8765/ping', '1'); } catch(e) {}
  }
  ping();
  setInterval(ping, 2000);
  window.addEventListener('beforeunload', function () {
    try { navigator.sendBeacon('http://127.0.0.1:8765/closed', '1'); } catch(e) {}
  });
})();
</script>
""", height=0)

# Keeps an open browser tab active by touching the Streamlit server periodically.
# This reduces idle sleep on many hosts, but a host-level hard restart still requires durable storage.
components.html("""
<script>
(function () {
  async function keepAwake(){
    try { await fetch(window.location.href, {cache: 'no-store'}); } catch(e) {}
  }
  setInterval(keepAwake, 5 * 60 * 1000);
})();
</script>
""", height=0)


def paths():
    return sorted([p for p in SCRAPED_DIR.glob("*.csv") if p.stat().st_size > 0])

@st.cache_data(show_spinner=False)
def load_data(path_strings, mtimes):
    return normalize_files([Path(p) for p in path_strings])

def clean_external_url(raw, auctioneer=""):
    """Return a real ad URL only. Suppresses Google search/map links and unwraps Google redirects."""
    url = str(raw or "").strip()
    if not url or url.lower() in {"nan", "none", "#"} or url.startswith("javascript:") or url.startswith("tw-ad:"):
        return ""
    if url.startswith("//"):
        url = "https:" + url
    if url.startswith("/"):
        base_by_auctioneer = {
            "TW": "https://www.tidewaterauctions.com",
            "AC": "https://realestate.alexcooper.com",
            "HW": "https://www.hwestauctions.com",
            "MWC": "https://apps.mwc-law.com",
            "BL": "https://ajbillig.com",
            "ADC": "https://www.auction.com",
        }
        base = base_by_auctioneer.get(str(auctioneer or "").upper(), "")
        url = base + url if base else ""

    # Unwrap google.com/url?q=<real auction ad> but do not show Google search/maps as Ad.
    for _ in range(2):
        try:
            parsed = urlparse(url)
            host = (parsed.netloc or "").lower()
            if host.endswith("google.com") or host.endswith("googleusercontent.com"):
                qs = parse_qs(parsed.query)
                target = (qs.get("q") or qs.get("url") or qs.get("u") or [""])[0]
                if target:
                    url = unquote(target)
                    continue
                return ""
        except Exception:
            return ""
        break

    low = url.lower()
    if not (low.startswith("http://") or low.startswith("https://")):
        return ""
    if "streamlit.app" in low or "localhost:8501" in low or "127.0.0.1:8501" in low:
        return ""
    if "google.com/search" in low or "google.com/maps" in low or "maps.google" in low:
        return ""
    return url

def export_view_df(df):
    """Flat investor-friendly export. Editable numbers stay in thousands."""
    out = pd.DataFrame()
    src = df.copy()
    dt = pd.to_datetime(src.get("Sale Date & Time"), errors="coerce")
    try:
        out["Date/Time"] = dt.dt.strftime("%-m/%-d/%Y %-I:%M %p").fillna("")
    except Exception:
        out["Date/Time"] = dt.dt.strftime("%m/%d/%Y %I:%M %p").fillna("")
    out["Address"] = src.get("Address", "")
    out["Auct"] = src.get("Auctioneer", "")
    out["Note"] = src.get("My Note", "")
    out["Deposit"] = src.get("Deposit", "")
    out["County"] = src.get("County", "")
    out["Occupied"] = src.get("Occupied", False)
    out["Look"] = src.get("Look", "")
    out["Comp"] = pd.to_numeric(src.get("Comp", 0), errors="coerce").fillna(0)
    out["Repair"] = pd.to_numeric(src.get("Rehab", 0), errors="coerce").fillna(0)
    out["Profit"] = pd.to_numeric(src.get("Profit", 0), errors="coerce").fillna(0)
    out["Max"] = ""
    out["%"] = ""
    out["MaxS"] = ""
    out["Ad"] = [clean_external_url(a, au) for a, au in zip(src.get("Ad Link", ""), src.get("Auctioneer", ""))]
    vals = [PROPERTY_VALUES.get(loose_address_key(a)) or {} for a in src.get("Address", "")]
    out["Est Value"] = [v.get("est") or "" for v in vals]
    out["MD Assessed"] = [v.get("assessed") or "" for v in vals]
    out["Sq Ft"] = [v.get("sqft") or "" for v in vals]
    out["Year Built"] = [v.get("year") or "" for v in vals]
    return out

def excel_bytes(df, sale_net=0.96, close1=0.06, close2=0.05):
    """Create a phone-friendly Excel export with live formulas.

    Inputs are in thousands, matching the auction workflow:
      Comp 700, Repair 100, Profit 100 -> Max = ((700*96%)-100-100)/(1+6%).
    User can edit Comp/Repair/Profit on the phone and Max/%/MaxS recalculate.
    """
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    bio = BytesIO()
    export_df = export_view_df(df)
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        export_df.to_excel(writer, index=False, sheet_name="Auctions")
        wb = writer.book
        ws = wb["Auctions"]
        ws.freeze_panes = "A2"

        # Header styling similar to Sam's auction spreadsheet.
        header_fill = PatternFill("solid", fgColor="4472C4")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D9E2F3")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Live formulas. Excel will recalculate when opened/edited.
        # I=Comp, J=Repair, K=Profit, L=Max, M=%, N=MaxS.
        for r in range(2, ws.max_row + 1):
            ws[f"L{r}"] = f'=IF(I{r}>0,(((I{r}*{sale_net})-J{r}-K{r})/(1+{close1}))*1000,"")'
            ws[f"M{r}"] = f'=IF(I{r}>0,L{r}/(I{r}*1000),"")'
            ws[f"N{r}"] = f'=IF(I{r}>0,(((I{r}*{sale_net})-J{r}-K{r})/(1+{close2}))*1000,"")'

            # Clickable Ad link without showing a long URL.
            ad_url = str(ws[f"O{r}"].value or "").strip()
            if ad_url:
                ws[f"O{r}"].hyperlink = ad_url
                ws[f"O{r}"].value = "Ad"
                ws[f"O{r}"].style = "Hyperlink"

        # Formatting.
        widths = {
            "A": 18, "B": 44, "C": 8, "D": 26, "E": 13, "F": 22, "G": 10, "H": 10,
            "I": 10, "J": 10, "K": 10, "L": 12, "M": 8, "N": 12, "O": 10,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for col in ["I", "J", "K", "L", "N"]:
                ws[f"{col}{row[0].row}"].number_format = '$#,##0'
            ws[f"M{row[0].row}"].number_format = '0%'

        # Make it sortable/filterable on phone/desktop.
        if ws.max_row >= 2:
            end_col = get_column_letter(ws.max_column)
            tab = Table(displayName="AuctionExport", ref=f"A1:{end_col}{ws.max_row}")
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            ws.add_table(tab)

        # Hidden settings sheet so the assumptions are preserved inside the file.
        settings = wb.create_sheet("Formula Settings")
        settings["A1"] = "Sale net"
        settings["B1"] = sale_net
        settings["A2"] = "Max closing cost"
        settings["B2"] = close1
        settings["A3"] = "MaxS closing cost"
        settings["B3"] = close2
        settings["A4"] = "Formula note"
        settings["B4"] = "Comp, Repair, and Profit are entered in thousands. Max and MaxS display as full dollar values."
        settings.sheet_state = "hidden"

        # Force recalculation in Excel/mobile apps.
        try:
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        except Exception:
            pass
    return bio.getvalue()

def safe_key(prefix, aid):
    return prefix + "_" + str(aid).replace("-", "_")

def remember_filter_state():
    keys = [
        "source_filter_sidebar", "hide_hidden_toggle", "hide_blocked_toggle",
        "auctioneer_grid_filter", "county_grid_filter", "search_grid_filter", "date_view_filter",
        "addr_w", "county_w", "note_w", "show_links_toggle", "show_ai_toggle"
    ]
    st.session_state["_restore_filters"] = {k: st.session_state.get(k) for k in keys if k in st.session_state}

def restore_filter_state():
    saved = st.session_state.pop("_restore_filters", None)
    if not saved:
        return
    for k, v in saved.items():
        st.session_state[k] = v

restore_filter_state()

def load_row_state(df, bids):
    if "loaded_ids" not in st.session_state:
        st.session_state.loaded_ids = set()
    latest = {}
    if not bids.empty and "Auction ID" in bids:
        latest = bids.sort_values("Saved At").drop_duplicates("Auction ID", keep="last").set_index("Auction ID").to_dict("index")
    for _, r in df.iterrows():
        aid = str(r["Auction ID"])
        if aid in st.session_state.loaded_ids:
            continue
        b = latest.get(aid, {})
        st.session_state[safe_key("occ", aid)] = bool(b.get("Occupied", r.get("Occupied", False)))
        st.session_state[safe_key("look", aid)] = str(b.get("Look", r.get("Look", "")) or "")
        def _blank_or_int(v):
            try:
                if v is None or str(v).strip().lower() in {"", "nan", "none"}:
                    return ""
                f = float(str(v).replace(",", ""))
                return "" if f == 0 else str(int(f))
            except Exception:
                return ""
        st.session_state[safe_key("comp", aid)] = _blank_or_int(b.get("Comp", r.get("Comp", "")))
        st.session_state[safe_key("rehab", aid)] = _blank_or_int(b.get("Rehab", r.get("Rehab", "")))
        st.session_state[safe_key("profit", aid)] = _blank_or_int(b.get("Profit", r.get("Profit", "")))
        note = b.get("My Note", r.get("My Note", "")) or ""
        if str(note).lower() in {"nan", "none"}:
            note = ""
        st.session_state[safe_key("note", aid)] = str(note)
        st.session_state.loaded_ids.add(aid)

def build_save_df(df, multiplier, sale_net, close1, close2):
    rows = []
    for _, r in df.iterrows():
        aid = str(r["Auction ID"])
        def _num(v):
            try:
                txt = str(v or "").replace(",", "").strip()
                return int(float(txt)) if txt else 0
            except Exception:
                return 0
        comp = _num(st.session_state.get(safe_key("comp", aid), ""))
        rehab = _num(st.session_state.get(safe_key("rehab", aid), ""))
        profit = _num(st.session_state.get(safe_key("profit", aid), ""))
        maxb, bidpct, maxs = calc_bid(comp, rehab, profit, sale_net, close1, close2, multiplier)
        row = r.to_dict()
        row.update({
            "Saved At": datetime.now().isoformat(timespec="seconds"),
            "Occupied": st.session_state.get(safe_key("occ", aid), False),
            "Look": st.session_state.get(safe_key("look", aid), ""),
            "Comp": comp, "Rehab": rehab, "Profit": profit,
            "Max": maxb, "%": bidpct, "MaxS": maxs,
            "My Note": st.session_state.get(safe_key("note", aid), ""),
        })
        rows.append(row)
    return pd.DataFrame(rows)

st.markdown('<div class="app-brand"><span class="badge">SIMO HOMES</span><span class="env">Maryland / DC foreclosure auctions</span></div>', unsafe_allow_html=True)
st.title("Auction Intelligence")
st.caption("Sources: Alex Cooper (AC) · Tidewater (TW) · Harvey West (HW) · McCabe Weisberg (MWC) · A. J. Billig (BL) · Auction.com in-person (ADC)")


def persist_user_state_before_refresh():
    """Save user-owned data before any scrape/cache refresh.

    Scraping is allowed to replace scraped CSV cache, but it must never wipe
    saved bids, notes, hidden properties, blocked cities, or layout choices.
    This also catches the common case where blocked city text was typed but
    the Save Blocked Cities button was not clicked before a refresh.
    """
    try:
        if "blocked_cities_text" in st.session_state:
            save_blocked_cities(set(str(st.session_state.get("blocked_cities_text", "")).splitlines()))
    except Exception:
        pass
    try:
        current = globals().get("filtered", pd.DataFrame())
        if isinstance(current, pd.DataFrame) and not current.empty:
            current_save = build_save_df(current, multiplier, sale_net, close1, close2)
            save_bids(pd.concat([load_bids(), current_save], ignore_index=True).sort_values("Saved At").drop_duplicates("Auction ID", keep="last"))
    except Exception:
        pass

with st.sidebar:
    st.header("Auction Sources")
    selected = st.multiselect("Sources", AUCTION_HOUSES, default=default_value("source_filter_sidebar", AUCTION_HOUSES), key="source_filter_sidebar")

    c1, c2 = st.columns(2)
    if c1.button("Scrape AC"):
        persist_user_state_before_refresh(); st.session_state.last_scrape = scrape_many(["AC"], clear_old=False); st.cache_data.clear(); st.session_state.pop("loaded_ids", None); st.rerun()
    if c2.button("Scrape TW"):
        persist_user_state_before_refresh(); st.session_state.last_scrape = scrape_many(["TW"], clear_old=False); st.cache_data.clear(); st.session_state.pop("loaded_ids", None); st.rerun()
    c3, c4 = st.columns(2)
    if c3.button("Scrape HW"):
        persist_user_state_before_refresh(); st.session_state.last_scrape = scrape_many(["HW"], clear_old=False); st.cache_data.clear(); st.session_state.pop("loaded_ids", None); st.rerun()
    if c4.button("Scrape MWC"):
        persist_user_state_before_refresh(); st.session_state.last_scrape = scrape_many(["MWC"], clear_old=False); st.cache_data.clear(); st.session_state.pop("loaded_ids", None); st.rerun()
    if st.button("Scrape BL (Howard Co. + more)", use_container_width=True):
        persist_user_state_before_refresh(); st.session_state.last_scrape = scrape_many(["BL"], clear_old=False); st.cache_data.clear(); st.session_state.pop("loaded_ids", None); st.rerun()
    if st.button("Scrape ADC (Auction.com)", use_container_width=True):
        persist_user_state_before_refresh(); st.session_state.last_scrape = scrape_many(["ADC"], clear_old=False); st.cache_data.clear(); st.session_state.pop("loaded_ids", None); st.rerun()

    if st.button("Full Refresh Selected", type="primary", use_container_width=True):
        persist_user_state_before_refresh(); st.session_state.last_scrape = scrape_many(selected, clear_old=True); st.cache_data.clear(); st.session_state.pop("loaded_ids", None); st.rerun()
    if st.button("Clear Scraped Cache", use_container_width=True):
        persist_user_state_before_refresh(); clear_cache(); st.cache_data.clear(); st.session_state.pop("loaded_ids", None); st.rerun()

    if "last_scrape" in st.session_state:
        st.write("Last scrape:")
        for s, res in st.session_state.last_scrape.items():
            if res["ok"]:
                st.success(f"{s}: {res['rows']} rows")
            else:
                st.error(f"{s}: {res['error']}")

    st.write(f"Cache files: **{len(paths())}**")
    st.caption("Storage: durable GitHub Gist backup is ON" if remote_enabled() else "Storage: local file only. On Streamlit Cloud, add GITHUB_TOKEN and GIST_ID secrets or archives can disappear after sleep/restart.")

    st.divider()
    st.header("Focus")
    hide_hidden = st.toggle("Hide hidden properties", default_value("hide_hidden_toggle", True), key="hide_hidden_toggle")
    hide_blocked = st.toggle("Hide blocked cities", default_value("hide_blocked_toggle", True), key="hide_blocked_toggle")
    favorites_only = st.toggle("Show favorited properties only", default_value("favorite_properties_only_toggle", False), key="favorite_properties_only_toggle")
    blocked = st.text_area("Blocked cities, one per line", value="\n".join(sorted(load_blocked_cities())), height=90, key="blocked_cities_text")
    if st.button("Save Blocked Cities", use_container_width=True):
        remember_filter_state()
        save_blocked_cities(set(st.session_state.get("blocked_cities_text", "").splitlines()))
        st.success("Blocked cities saved. Current filters kept.")
        st.rerun()
    if st.button("Clear Hidden Properties", use_container_width=True):
        clear_hidden(); st.rerun()

    st.divider()
    st.header("Formula")
    thousands = st.toggle("Inputs are in thousands", True)
    multiplier = 1000 if thousands else 1
    sale_net = st.number_input("Sale net (%)", 0.0, 100.0, 96.0, .5) / 100
    close1 = st.number_input("Max closing cost (%)", 0.0, 25.0, 6.0, .25) / 100
    close2 = st.number_input("MaxS closing cost (%)", 0.0, 25.0, 5.0, .25) / 100

    st.divider()
    st.header("Optional")
    show_ai = st.toggle("Show Sam AI columns", default_value("show_ai_toggle", False), key="show_ai_toggle")
    show_links = st.toggle("Show Zillow/Redfin links", default_value("show_links_toggle", True), key="show_links_toggle")

    st.divider()
    st.header("Column Widths")
    st.caption("Adjust grid column widths here. Streamlit does not support true drag-resize for this custom editable row layout.")
    addr_w = st.slider("Address width", 1.5, 5.0, float(default_value("addr_w", 2.3)), 0.25, key="addr_w")
    county_w = st.slider("County width", 0.7, 2.5, float(default_value("county_w", 1.0)), 0.1, key="county_w")
    note_w = st.slider("Note width", 0.8, 3.0, float(default_value("note_w", 1.0)), 0.1, key="note_w")

    if st.button("Save Default Layout", use_container_width=True):
        save_layout_defaults({
            "source_filter_sidebar": st.session_state.get("source_filter_sidebar", AUCTION_HOUSES),
            "favorite_properties_only_toggle": st.session_state.get("favorite_properties_only_toggle", False),
            "auctioneer_grid_filter": st.session_state.get("auctioneer_grid_filter", []),
            "county_grid_filter": st.session_state.get("county_grid_filter", []),
            "date_view_filter": st.session_state.get("date_view_filter", "Current auction week"),
            "hide_hidden_toggle": st.session_state.get("hide_hidden_toggle", True),
            "hide_blocked_toggle": st.session_state.get("hide_blocked_toggle", True),
            "show_links_toggle": st.session_state.get("show_links_toggle", True),
            "show_ai_toggle": st.session_state.get("show_ai_toggle", False),
            "addr_w": st.session_state.get("addr_w", 2.5),
            "county_w": st.session_state.get("county_w", 1.0),
            "note_w": st.session_state.get("note_w", 1.2),
        })
        st.success("Default layout saved.")


p = paths()
# IMPORTANT: do not auto-run a full scrape on app launch.
# The launcher is portable and may be opened on a new computer; forcing AC/TW/HW/MWC
# scraping immediately can make startup look frozen for many minutes.
# Load cached CSVs instantly, then let the user refresh selected sources manually.
raw = load_data([str(x) for x in p], [x.stat().st_mtime for x in p]) if p else pd.DataFrame()
bids = load_bids()
df = merge_bids(raw, bids) if not raw.empty else raw

if not df.empty:
    df["County"] = df["County"].apply(lambda x: x if x in MD_COUNTIES else "Unknown County")

def _loose_addr_key(addr):
    """Street number + first real street word + zip. Catches the same property listed by
    two sources with slightly different formatting (e.g. '5865E BONIWOOD TURN' vs '5865 E Boniwood Turn')."""
    a = str(addr or "").upper()
    num = re.match(r"\s*(\d+)", a)
    num = num.group(1) if num else ""
    words = [w for w in re.findall(r"[A-Z]+", a.split(",")[0]) if w not in {"E", "W", "N", "S", "NE", "NW", "SE", "SW", "UNIT", "APT", "STE"}]
    zipm = re.search(r"\b(\d{5})(?:-\d{4})?\s*$", a)
    return (num + "|" + (words[0] if words else "") + "|" + (zipm.group(1) if zipm else "")) if num else a

def dedupe_across_sources(frame):
    """One row per property across ALL sources. Prefer the auctioneer that actually runs
    the sale (AC/TW/HW/MWC/BL) over a listing-only source (ADC), then prefer rows with an ad."""
    if frame.empty or "Address" not in frame:
        return frame
    pref = {"ADC": 9}
    f = frame.copy()
    f["_lk"] = f["Address"].apply(_loose_addr_key)
    f["_rank"] = f["Auctioneer"].map(lambda x: pref.get(str(x).upper(), 0)) + f["Ad Link"].apply(lambda x: 0 if str(x or "").strip() and str(x).lower() not in {"nan", "none"} else 1)
    f = f.sort_values(["_lk", "_rank"]).drop_duplicates("_lk", keep="first")
    return f.drop(columns=["_lk", "_rank"])

if not df.empty:
    df = dedupe_across_sources(df)

# Favorite properties are user preferences, not scrape data. Keep them
# separate so refreshes cannot wipe them.
favorite_properties = load_favorite_properties()
if st.session_state.get("favorite_properties_only_toggle", False) and not df.empty:
    df = df[df["Auction ID"].astype(str).isin(favorite_properties)]

if df.empty:
    st.warning("No cached auction data found. Select sources in the sidebar and click a scrape/refresh button. First setup on a new computer installs dependencies once, but scraping is manual so startup stays fast.")
    st.stop()

# focus filters
if hide_hidden:
    hidden = load_hidden()
    if hidden:
        df = df[~df["Address"].apply(lambda x: address_key(x) in hidden)]
if hide_blocked:
    blocked_set = load_blocked_cities()
    if blocked_set:
        df = df[~df["Address"].apply(lambda x: city_from_address(x) in blocked_set)]


# ---------------------------------------------------------------------------
# Phone / field mode. Default comes from the browser's User-Agent (st.context.headers),
# ?view=phone|desktop overrides it (bookmark /?view=phone on the iPhone), and the
# sidebar / top toggle lets Sam switch any time.
# ---------------------------------------------------------------------------
def _ua_is_mobile():
    try:
        ua = str(st.context.headers.get("User-Agent", "") or "")
    except Exception:
        ua = ""
    return bool(re.search(r"iPhone|Android.*Mobile|Mobile Safari|iPod|Windows Phone", ua, re.I))

_qp_view = str(st.query_params.get("view", "") or "").lower()
if _qp_view in {"phone", "desktop"}:
    st.session_state["view_mode"] = _qp_view
elif "view_mode" not in st.session_state:
    st.session_state["view_mode"] = "phone" if _ua_is_mobile() else "desktop"
PHONE = st.session_state["view_mode"] == "phone"

with st.sidebar:
    st.divider()
    st.header("Layout")
    _vm = st.radio("Layout", ["Desktop grid", "Phone cards"], index={"desktop": 0, "phone": 1}[st.session_state["view_mode"]], horizontal=True, key="view_mode_radio", label_visibility="collapsed")
    _new_mode = {"Desktop grid": "desktop", "Phone cards": "phone"}[_vm]
    if _new_mode != st.session_state["view_mode"]:
        st.session_state["view_mode"] = _new_mode
        st.query_params["view"] = _new_mode
        st.rerun()

def tw_ad_html(link):
    """Return cached Tidewater ad HTML for a 'tw-ad:<id>' link, or ''."""
    try:
        ad_id = str(link).split(":", 1)[1].strip()
        path = SCRAPED_DIR / "ads" / f"TW_{ad_id}.html"
        if path.exists():
            return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        pass
    return ""

@st.dialog("Legal ad", width="large")
def show_ad_dialog(html):
    st.markdown('<div class="ad-body">' + html + "</div>", unsafe_allow_html=True)

def render_ad_control(container, raw_link, auctioneer, aid, label="Ad", compact=True):
    """Ad button: external link pill for real URLs; in-app viewer for Tidewater (tw-ad:)."""
    raw = str(raw_link or "").strip()
    if raw.startswith("tw-ad:"):
        html = tw_ad_html(raw)
        if html:
            if container.button(label, key=safe_key("ad", aid), use_container_width=True, help="View Tidewater legal ad"):
                show_ad_dialog(html)
            return True
        container.markdown(_cell("—", "muted"), unsafe_allow_html=True)
        return False
    link = clean_external_url(raw, auctioneer)
    if link:
        container.markdown(f'<a class="pill{"" if compact else " wide"}" href="{link}" target="_blank" rel="noopener noreferrer">{label} ↗</a>', unsafe_allow_html=True)
        return True
    container.markdown(_cell("—", "muted"), unsafe_allow_html=True)
    return False

def _cell(html, cls=""):
    return f'<div class="gc {cls}">{html}</div>'

PROPERTY_VALUES = load_values()

def _k(n):
    try:
        n = float(n)
    except Exception:
        return ""
    if not n:
        return ""
    return f"${n/1000:,.0f}k" if n >= 10000 else f"${n:,.0f}"

def value_info(address):
    """(headline, label, tooltip, record) for a property: market estimate if we have one, else MD assessment."""
    rec = PROPERTY_VALUES.get(loose_address_key(address)) or {}
    est, assd = rec.get("est"), rec.get("assessed")
    head, label = "", ""
    if est:
        head, label = _k(est), "est"
    elif assd:
        head, label = _k(assd), "assd"
    bits = []
    if est:
        bits.append(f"Est. market value {money(est)}" + (f" (range {_k(rec.get('est_low'))}–{_k(rec.get('est_high'))})" if rec.get("est_low") else ""))
    if assd:
        bits.append(f"MD assessment {money(assd)}")
    if rec.get("rent"):
        bits.append(f"Est. rent {money(rec['rent'])}/mo")
    if rec.get("sqft"):
        bits.append(f"{int(rec['sqft']):,} sq ft")
    if rec.get("year"):
        bits.append(f"built {int(rec['year'])}")
    if rec.get("last_sale"):
        bits.append(f"last sale {money(rec['last_sale'])} {str(rec.get('last_sale_date') or '')[:4]}".strip())
    if rec.get("owner_occupied"):
        bits.append("owner-occupied")
    return head, label, " · ".join(bits), rec

# ---------------------------------------------------------------------------
# Filters
# ---------------------------------------------------------------------------
for _lk, _lv in {
    "auctioneer_grid_filter": default_value("auctioneer_grid_filter", []),
    "county_grid_filter": default_value("county_grid_filter", DEFAULT_COUNTY_FILTER),
    "date_view_filter": default_value("date_view_filter", "Current auction week"),
}.items():
    if _lk not in st.session_state:
        st.session_state[_lk] = _lv
county_values = set(df["County"].dropna().astype(str))
county_options = [c for c in MD_COUNTIES if c in county_values]
county_default = [c for c in default_value("county_grid_filter", DEFAULT_COUNTY_FILTER) if c in county_options]
if "county_grid_filter" not in st.session_state:
    st.session_state["county_grid_filter"] = county_default

if PHONE:
    with st.expander("Filters", expanded=False):
        search = st.text_input("Search", key="search_grid_filter", placeholder="Address, city, county…")
        date_view = st.radio("Date view", ["Current auction week", "All future", "All dates"], horizontal=True, key="date_view_filter")
        county_filter = st.multiselect("County", county_options, key="county_grid_filter")
        auctioneer_filter = st.multiselect("Auctioneer", sorted(df["Auctioneer"].dropna().unique()), key="auctioneer_grid_filter")
else:
    st.subheader("Filters")
    with st.container(border=True):
        f1, f2, f3, f4 = st.columns([1, 1.6, 1.8, 1.2])
        auctioneer_filter = f1.multiselect("Auctioneer", sorted(df["Auctioneer"].dropna().unique()), key="auctioneer_grid_filter")
        county_filter = f2.multiselect("County", county_options, key="county_grid_filter")
        search = f3.text_input("Search", key="search_grid_filter")
        date_view = f4.radio("Date view", ["Current auction week", "All future", "All dates"], horizontal=False, key="date_view_filter")

filtered = df.copy()
if auctioneer_filter:
    filtered = filtered[filtered["Auctioneer"].isin(auctioneer_filter)]
if county_filter:
    filtered = filtered[filtered["County"].isin(county_filter)]
if search.strip():
    q = search.lower().strip()
    filtered = filtered[(filtered["Address"] + " " + filtered["County"] + " " + filtered["Deposit"]).str.lower().str.contains(q, na=False)]

dates = pd.to_datetime(filtered["Sale Date & Time"], errors="coerce").dt.date
if date_view == "Current auction week":
    a, b = this_or_next_week()
    filtered = filtered[(dates >= a) & (dates <= b)]
elif date_view == "All future":
    from datetime import date
    filtered = filtered[dates >= date.today()]

if favorite_properties and not filtered.empty:
    filtered["_FavProperty"] = filtered["Auction ID"].astype(str).isin(favorite_properties)
    filtered = filtered.sort_values(["_FavProperty", "Sale Date & Time", "County", "Address"], ascending=[False, True, True, True]).drop(columns=["_FavProperty"], errors="ignore")
else:
    filtered = filtered.sort_values(["Sale Date & Time", "County", "Address"])
if filtered.empty:
    st.info("No auctions match filters.")
    st.stop()

load_row_state(filtered, bids)
visible_save = build_save_df(filtered, multiplier, sale_net, close1, close2)
if not visible_save.empty:
    existing = load_bids()
    combined = pd.concat([existing, visible_save], ignore_index=True)
    combined = combined.sort_values("Saved At").drop_duplicates("Auction ID", keep="last")
    save_bids(combined)

# ---------------------------------------------------------------------------
# Actions
# ---------------------------------------------------------------------------
if PHONE:
    with st.expander("Save / Export / Backup", expanded=False):
        if st.button("Save Bids", type="primary", use_container_width=True):
            save_bids(pd.concat([load_bids(), visible_save], ignore_index=True).sort_values("Saved At").drop_duplicates("Auction ID", keep="last"))
            st.success("Saved")
        st.download_button("Export Excel", data=excel_bytes(visible_save, sale_net, close1, close2), file_name="auction_intelligence.xlsx", use_container_width=True)
        st.download_button("Backup Archive", data=load_bids().to_csv(index=False), file_name="auction_archive_backup.csv", mime="text/csv", use_container_width=True)
else:
    st.subheader("Actions")
    top1, top2, top3, top4 = st.columns([1,1,1,3])
    if top1.button("Save Bids", type="primary", use_container_width=True):
        save_bids(pd.concat([load_bids(), visible_save], ignore_index=True).sort_values("Saved At").drop_duplicates("Auction ID", keep="last"))
        st.success("Saved")
    top2.download_button("Export Excel", data=excel_bytes(visible_save, sale_net, close1, close2), file_name="auction_intelligence.xlsx", use_container_width=True)
    current_archive_csv = load_bids().to_csv(index=False)
    top3.download_button("Backup Archive", data=current_archive_csv, file_name="auction_archive_backup.csv", mime="text/csv", use_container_width=True)
    with top4.expander("Restore archive from backup CSV"):
        restore_file = st.file_uploader("Upload auction_archive_backup.csv", type=["csv"], key="restore_archive_csv")
        if restore_file is not None and st.button("Restore Archive", type="primary"):
            try:
                restored = pd.read_csv(restore_file)
                save_bids(pd.concat([load_bids(), restored], ignore_index=True).sort_values("Saved At").drop_duplicates("Auction ID", keep="last"))
                st.success("Archive restored and saved.")
                st.rerun()
            except Exception as e:
                st.error(f"Could not restore CSV: {e}")

# ---------------------------------------------------------------------------
# Auction grid / cards
# ---------------------------------------------------------------------------
st.subheader("Auction Grid" if not PHONE else "Auctions")
if date_view == "Current auction week":
    a, b = this_or_next_week()
    st.caption(f"Auction week {a.strftime('%b %d')} – {b.strftime('%b %d, %Y')} · {len(filtered)} active sale{'s' if len(filtered) != 1 else ''}. Comp / Rehab / Profit are in thousands.")
else:
    st.caption(f"{len(filtered)} active sale{'s' if len(filtered) != 1 else ''}. Comp / Rehab / Profit are in thousands.")

filtered["_Date"] = pd.to_datetime(filtered["Sale Date & Time"], errors="coerce").dt.date

def _to_int(v):
    try:
        txt = str(v or "").replace(",", "").strip()
        return int(float(txt)) if txt else 0
    except Exception:
        return 0

def row_numbers(aid):
    comp = _to_int(st.session_state.get(safe_key("comp", aid), ""))
    rehab = _to_int(st.session_state.get(safe_key("rehab", aid), ""))
    profit = _to_int(st.session_state.get(safe_key("profit", aid), ""))
    maxb, bidpct, maxs = calc_bid(comp, rehab, profit, sale_net, close1, close2, multiplier)
    return comp, rehab, profit, maxb, bidpct, maxs

def render_phone_card(r, gi):
    aid = str(r["Auction ID"])
    dt = pd.to_datetime(r["Sale Date & Time"], errors="coerce")
    time_txt = "" if pd.isna(dt) else dt.strftime("%I:%M %p").lstrip("0")
    src = str(r["Auctioneer"]).upper()
    is_fav = aid in favorite_properties
    county_txt = str(r["County"]).replace(" County", "")
    vh, vl, vtip, vrec = value_info(r["Address"])
    with st.container(border=True, key=f"pcard_{aid}"):
        h1, h2 = st.columns([6, 1])
        h1.markdown(
            f'<div class="pc-head"><div class="pc-top"><span class="pc-time">{time_txt or "Time TBD"}</span>'
            f'<span class="src {src}">{src}</span><span class="pc-county">{county_txt}</span></div>'
            f'<div class="pc-addr">{r["Address"]}</div>'
            f'<div class="pc-meta">Deposit <b>{r["Deposit"]}</b>'
            + (f' · <b>{r.get("Occupancy","")}</b>' if str(r.get("Occupancy","") or "").strip() else "")
            + '</div>'
            + (f'<div class="pc-val"><b>{vh}</b> <span class="vl">{ {"est": "est. value", "assd": "MD assessed"}.get(vl, vl) }</span>' + (f' <span class="more">· {vtip.split(" · ", 1)[1] if " · " in vtip else ""}</span>' if " · " in vtip else "") + '</div>' if vh else "")
            + '</div>',
            unsafe_allow_html=True,
        )
        if h2.button("", icon=":material/star:" if is_fav else ":material/star_outline:", key=safe_key("fav", aid), type="primary" if is_fav else "secondary", use_container_width=True):
            toggle_favorite_property(aid)
            st.rerun()

        i1, i2, i3 = st.columns(3)
        i1.text_input("Comp (k)", key=safe_key("comp", aid), placeholder="Comp")
        i2.text_input("Rehab (k)", key=safe_key("rehab", aid), placeholder="Rehab")
        i3.text_input("Profit (k)", key=safe_key("profit", aid), placeholder="Profit")
        comp, rehab, profit, maxb, bidpct, maxs = row_numbers(aid)
        st.markdown(
            f'<div class="pc-bids"><div><span class="k">MAX BID</span><span class="v">{money(maxb) or "—"}</span></div>'
            f'<div><span class="k">%</span><span class="v">{pct(bidpct) or "—"}</span></div>'
            f'<div><span class="k">MAX S</span><span class="v">{money(maxs) or "—"}</span></div></div>',
            unsafe_allow_html=True,
        )
        j1, j2, j3 = st.columns([1.1, 1, 2.2])
        j1.selectbox("Look", ["", "Y", "N", "YY", "Soso"], key=safe_key("look", aid), placeholder="Look")
        j2.checkbox("Occupied", key=safe_key("occ", aid))
        j3.text_input("Note", key=safe_key("note", aid), placeholder="Note")
        sdat_url = (vrec or {}).get("sdat_url", "")
        b1, b2, b3, b4, b5 = st.columns([1, 1, 1, 1, .6])
        render_ad_control(b1, r.get("Ad Link", ""), r.get("Auctioneer", ""), aid, label="Ad", compact=False)
        b2.markdown(f'<a class="pill wide" href="{zillow_link(r["Address"])}" target="_blank" rel="noopener noreferrer">Zillow</a>', unsafe_allow_html=True)
        b3.markdown(f'<a class="pill wide" href="{redfin_search_link(r["Address"])}" target="_blank" rel="noopener noreferrer">Redfin</a>', unsafe_allow_html=True)
        b4.markdown(f'<a class="pill wide" href="{sdat_url}" target="_blank" rel="noopener noreferrer">SDAT</a>' if sdat_url else '<span class="pill wide muted">SDAT</span>', unsafe_allow_html=True)
        if b5.button("", icon=":material/visibility_off:", key=safe_key("hide", aid), help="Hide this property", use_container_width=True):
            hide_address(r["Address"])
            st.rerun()

for gi, (d, group) in enumerate(filtered.groupby("_Date", dropna=False)):
    day_label = pd.to_datetime(d).strftime("%A, %B %d, %Y") if pd.notna(d) else "Unknown Date"
    n = len(group)
    if PHONE:
        st.markdown(f'<div class="day-header phone"><span class="d">{pd.to_datetime(d).strftime("%a, %b %d") if pd.notna(d) else "Unknown Date"}</span><span class="n">{n} sale{"s" if n != 1 else ""}</span></div>', unsafe_allow_html=True)
        for _, r in group.iterrows():
            render_phone_card(r, gi)
        continue

    with st.container(border=True, key=f"grid_card_{gi}"):
        st.markdown(f'<div class="day-header"><span class="d">{day_label}</span><span class="n">{n} sale{"s" if n != 1 else ""}</span></div>', unsafe_allow_html=True)
        widths = [.42,.9,.55,county_w,addr_w,.8,.85,.5,.78,.7,.7,.7,.85,.55,.85,note_w,.6]
        headers = ["", "Time", "Src", "County", "Address", "Dep", "Value", "Occ", "Look", "Comp", "Rehab", "Profit", "Max", "%", "MaxS", "Note", "Ad"]
        if show_ai:
            widths += [.9,.9]
            headers += ["AI ARV","AI Max"]
        if show_links:
            widths += [1.1]
            headers += ["Links"]
        widths += [.45]
        headers += [""]
        cols = st.columns(widths)
        for c, h in zip(cols, headers):
            c.markdown(f'<div class="gh">{h}&nbsp;</div>', unsafe_allow_html=True)

        for _, r in group.iterrows():
            aid = str(r["Auction ID"])
            cols = st.columns(widths)
            dt = pd.to_datetime(r["Sale Date & Time"], errors="coerce")
            is_fav = aid in favorite_properties
            if cols[0].button("", icon=":material/star:" if is_fav else ":material/star_outline:", key=safe_key("fav", aid), help="Favorite", use_container_width=True, type="primary" if is_fav else "secondary"):
                toggle_favorite_property(aid)
                st.rerun()
            time_txt = "" if pd.isna(dt) else dt.strftime("%I:%M %p").lstrip("0")
            cols[1].markdown(_cell(time_txt, "time strong"), unsafe_allow_html=True)
            src = str(r["Auctioneer"]).upper()
            cols[2].markdown(f'<span class="src {src}">{src}</span>', unsafe_allow_html=True)
            county_txt = str(r["County"]).replace(" County", "")
            cols[3].markdown(_cell(county_txt, "muted county"), unsafe_allow_html=True)
            cols[4].markdown(_cell(r["Address"], "addr"), unsafe_allow_html=True)
            cols[5].markdown(_cell(r["Deposit"], "mono"), unsafe_allow_html=True)
            vh, vl, vtip, vrec = value_info(r["Address"])
            if vh:
                cols[6].markdown(f'<div class="gc mono val" title="{vtip}"><span class="vh">{vh}</span><span class="vl">{vl}</span></div>', unsafe_allow_html=True)
            else:
                cols[6].markdown(_cell("—", "muted"), unsafe_allow_html=True)
            cols[7].checkbox("Occ", key=safe_key("occ", aid), label_visibility="collapsed")
            cols[8].selectbox("Look", ["", "Y", "N", "YY", "Soso"], key=safe_key("look", aid), label_visibility="collapsed", placeholder="—")
            cols[9].text_input("Comp", key=safe_key("comp", aid), label_visibility="collapsed", placeholder="—")
            cols[10].text_input("Rehab", key=safe_key("rehab", aid), label_visibility="collapsed", placeholder="—")
            cols[11].text_input("Profit", key=safe_key("profit", aid), label_visibility="collapsed", placeholder="—")
            comp, rehab, profit, maxb, bidpct, maxs = row_numbers(aid)
            cols[12].markdown(_cell(money(maxb) or "—", "mono strong" if maxb else "mono muted"), unsafe_allow_html=True)
            cols[13].markdown(_cell(pct(bidpct) or "—", "mono" if bidpct else "mono muted"), unsafe_allow_html=True)
            cols[14].markdown(_cell(money(maxs) or "—", "mono" if maxs else "mono muted"), unsafe_allow_html=True)
            cols[15].text_input("Note", key=safe_key("note", aid), label_visibility="collapsed", placeholder="Note")
            render_ad_control(cols[16], r.get("Ad Link", ""), r.get("Auctioneer", ""), aid)
            idx = 17
            if show_ai:
                cols[idx].markdown(_cell(money(float(comp or 0) * multiplier) if comp else "—", "mono"), unsafe_allow_html=True)
                cols[idx+1].markdown(_cell(money(maxb) or "—", "mono"), unsafe_allow_html=True)
                idx += 2
            if show_links:
                addr = r["Address"]
                sdat_url = (vrec or {}).get("sdat_url", "")
                sdat_a = f'<a href="{sdat_url}" target="_blank" rel="noopener noreferrer" title="MD SDAT record">S</a>' if sdat_url else ""
                cols[idx].markdown(f'<div class="links"><a href="{zillow_link(addr)}" target="_blank" rel="noopener noreferrer" title="Zillow">Z</a><a href="{redfin_search_link(addr)}" target="_blank" rel="noopener noreferrer" title="Redfin">R</a>{sdat_a}</div>', unsafe_allow_html=True)
                idx += 1
            if cols[idx].button("", icon=":material/visibility_off:", key=safe_key("hide", aid), help="Hide this property", use_container_width=True):
                hide_address(r["Address"])
                st.rerun()

st.divider()
tab1, tab2 = st.tabs(["Saved Bid Archive", "AI Coach"])
with tab1:
    arch = load_bids()
    if arch.empty:
        st.info("No saved bids.")
    else:
        arch_view = arch.drop(columns=["Zestimate", "Redfin Estimate"], errors="ignore")
        st.dataframe(arch_view.sort_values("Saved At", ascending=False), hide_index=True, use_container_width=True)
with tab2:
    arch = load_bids()
    if arch.empty:
        st.info("Save bids first.")
    else:
        comp = pd.to_numeric(arch["Comp"], errors="coerce").fillna(0)
        maxb = pd.to_numeric(arch["Max"], errors="coerce").fillna(0)
        valid = comp > 0
        st.metric("Bids Learned", len(arch))
        st.metric("Average Bid %", pct((maxb[valid] / (comp[valid] * multiplier)).mean()) if valid.any() else "")
